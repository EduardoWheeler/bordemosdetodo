import os
from docx import Document
from docx.shared import Pt, RGBColor

def exportar_codigo_a_word(directorio_raiz, archivo_salida="Codigo_Proyecto.docx"):
    doc = Document()
    doc.add_heading('Documentación de Código Fuente', level=0)

    # Carpetas o archivos que quieras ignorar para que no sature el documento
    ignorar_carpetas = {'.git', '__pycache__', 'venv', '.venv', 'build', 'dist', '.idea'}

    print("Procesando archivos...")

    for root, dirs, files in os.walk(directorio_raiz):
        # Filtrar carpetas ignoradas
        dirs[:] = [d for d in dirs if d not in ignorar_carpetas]

        for file in files:
            # Puedes agregar o quitar extensiones según lo que necesites
            if file.endswith(('.py', '.sql', '.html', '.css', '.js', '.json')):
                ruta_completa = os.path.join(root, file)
                ruta_relativa = os.path.relpath(ruta_completa, directorio_raiz)

                print(f"Agregando: {ruta_relativa}")

                # Crear un título en el Word con la ruta del archivo
                doc.add_heading(f"Archivo: {ruta_relativa}", level=2)

                try:
                    with open(ruta_completa, 'r', encoding='utf-8') as f:
                        contenido = f.read()

                    # Agregar el código formateado con tipografía de código
                    p = doc.add_paragraph()
                    run = p.add_run(contenido)
                    run.font.name = 'Consolas'  # Tipografía monoespaciada
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(40, 40, 40)

                except Exception as e:
                    doc.add_paragraph(f"[Error al leer el archivo: {e}]")

    doc.save(archivo_salida)
    print(f"\n¡Listo! Documento creado exitosamente como: {archivo_salida}")

if __name__ == '__main__':
    # El punto '.' significa la carpeta actual donde estás ejecutando el script
    exportar_codigo_a_word('.')